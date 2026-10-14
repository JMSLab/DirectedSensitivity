import os
import re
import sys

import numpy as np
import openpyxl
import scipy.optimize
import statsmodels.api as sm

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import config
import specification as spec

TARGETS_PATH = config.PATHS['targets']
RESULTS_PATH = config.output('replicate.txt')

ASSERTION_TOLERANCE = 0.05
PARAMETERS = ('demand_intercept', 'demand_slope', 'avg_cost_intercept', 'avg_cost_slope')


def Main():
    os.makedirs(os.path.dirname(RESULTS_PATH), exist_ok=True)
    outfile = open(RESULTS_PATH, 'w')
    sys.stdout = _Tee(sys.__stdout__, outfile)

    data    = spec.DATA
    targets = load_targets(TARGETS_PATH)
    ols_est = fit_ols(data)
    gmm_est = fit_gmm(data, ols_est)
    welfare = compute_welfare(gmm_est)
    print_panel_a(ols_est, gmm_est, targets)
    print_panel_b(welfare, targets)
    check_assertions(ols_est, gmm_est, welfare, targets, ASSERTION_TOLERANCE)


class _Tee:
    def __init__(self, *files):
        self._files = files

    def write(self, text):
        for f in self._files:
            f.write(text)

    def flush(self):
        for f in self._files:
            f.flush()


def load_targets(path):
    ws = openpyxl.load_workbook(path).active
    return {
        'demand_intercept':   _parse_num(ws['B8'].value),
        'demand_slope':       _parse_num(ws['B5'].value),
        'avg_cost_intercept': _parse_num(ws['C8'].value),
        'avg_cost_slope':     _parse_num(ws['C5'].value),
        # Notation follows EFC2010 (Figure V)
        'Q_C':   _parse_qp(ws['C15'].value)[0],
        'P_C':   _parse_qp(ws['C15'].value)[1],
        'Q_E':   _parse_qp(ws['C16'].value)[0],
        'P_E':   _parse_qp(ws['C16'].value)[1],
        'CDE':   _parse_num(ws['C17'].value),
        'ABE':   _parse_num(ws['C18'].value),
        'EGH':   _parse_num(ws['C19'].value),
    }


def _parse_num(val):
    if isinstance(val, (int, float)):
        return float(val)
    return float(str(val).strip().replace('−', '-'))


def _parse_qp(val):
    nums = re.findall(r'[\d.]+', str(val))
    return float(nums[0]), float(nums[1])


def fit_ols(data):
    demand_intercept, demand_slope     = spec.FitLine(data['p'], data['D'], data['w1'])
    avg_cost_intercept, avg_cost_slope = spec.FitLine(data['p'], data['c'], data['w2'])
    return {'demand_intercept': demand_intercept, 'demand_slope': demand_slope,
            'avg_cost_intercept': avg_cost_intercept, 'avg_cost_slope': avg_cost_slope}


def fit_gmm(data, ols_est):
    p, D, c, w1, w2 = data['p'], data['D'], data['c'], data['w1'], data['w2']
    X = sm.add_constant(p)

    def gmm_obj(y, w):
        W = np.diag(w / w.sum())
        def objective(theta):
            resid = y - X @ theta
            return float(resid @ W @ resid)
        return objective

    options = {'xatol': config.CONFIG['tolerances']['optimizer'],
               'fatol': config.CONFIG['tolerances']['optimizer'],
               'maxiter': config.CONFIG['iterations']['optimizer']}
    res_demand = scipy.optimize.minimize(
        gmm_obj(D, w1), [ols_est['demand_intercept'], ols_est['demand_slope']],
        method='Nelder-Mead', options=options,
    )
    res_cost = scipy.optimize.minimize(
        gmm_obj(c, w2), [ols_est['avg_cost_intercept'], ols_est['avg_cost_slope']],
        method='Nelder-Mead', options=options,
    )
    demand_intercept, demand_slope     = res_demand.x
    avg_cost_intercept, avg_cost_slope = res_cost.x
    return {'demand_intercept': demand_intercept, 'demand_slope': demand_slope,
            'avg_cost_intercept': avg_cost_intercept, 'avg_cost_slope': avg_cost_slope}


def linear_inv_demand(Q, est):
    return (Q - est['demand_intercept']) / est['demand_slope']


def linear_ac(Q, est):
    return est['avg_cost_intercept'] + est['avg_cost_slope'] * linear_inv_demand(Q, est)


def linear_mc(Q, est):
    demand_intercept, demand_slope = est['demand_intercept'], est['demand_slope']
    avg_cost_intercept, avg_cost_slope = est['avg_cost_intercept'], est['avg_cost_slope']
    return (avg_cost_intercept - demand_intercept * avg_cost_slope / demand_slope
            + 2 * (avg_cost_slope / demand_slope) * Q)


def linear_equilibria(est):
    demand_intercept, demand_slope = est['demand_intercept'], est['demand_slope']
    avg_cost_intercept, avg_cost_slope = est['avg_cost_intercept'], est['avg_cost_slope']
    Q_C = demand_intercept + demand_slope * avg_cost_intercept / (1 - avg_cost_slope)
    Q_E = ((demand_slope * avg_cost_intercept + demand_intercept * (1 - avg_cost_slope))
           / (1 - 2 * avg_cost_slope))
    return {'Q_C': Q_C, 'P_C': linear_inv_demand(Q_C, est),
            'Q_E': Q_E, 'P_E': linear_inv_demand(Q_E, est)}


def compute_welfare(gmm_est):
    eq = linear_equilibria(gmm_est)
    Q_C, Q_E = eq['Q_C'], eq['Q_E']
    CDE = 0.5 * (Q_E - Q_C) * (eq['P_C'] - linear_mc(Q_C, gmm_est))
    ABE = 0.5 * Q_E * (linear_inv_demand(0, gmm_est) - linear_mc(0, gmm_est))
    EGH = 0.5 * (1 - Q_E) * (linear_mc(1, gmm_est) - linear_inv_demand(1, gmm_est))
    return {**eq, 'CDE': CDE, 'ABE': ABE, 'EGH': EGH}


def print_panel_a(ols_est, gmm_est, targets):
    print('Panel A: OLS')
    print(f"  Demand  intercept={ols_est['demand_intercept']:.2f}  "
          f"slope={ols_est['demand_slope']:.5f}  "
          f"(targets intercept={targets['demand_intercept']}  slope={targets['demand_slope']})")
    print(f"  Cost    intercept={ols_est['avg_cost_intercept']:.2f}  "
          f"slope={ols_est['avg_cost_slope']:.5f}  "
          f"(targets intercept={targets['avg_cost_intercept']}  slope={targets['avg_cost_slope']})")
    print('\nPanel A: GMM')
    print(f"  Demand  intercept={gmm_est['demand_intercept']:.2f}  "
          f"slope={gmm_est['demand_slope']:.5f}")
    print(f"  Cost    intercept={gmm_est['avg_cost_intercept']:.2f}  "
          f"slope={gmm_est['avg_cost_slope']:.5f}")


def print_panel_b(welfare, targets):
    print('\nPanel B:')
    print(f"  Competitive  Q_C={welfare['Q_C']:.3f} (target {targets['Q_C']})  "
          f"P_C={welfare['P_C']:.1f} (target {targets['P_C']})")
    print(f"  Efficient    Q_E={welfare['Q_E']:.3f} (target {targets['Q_E']})  "
          f"P_E={welfare['P_E']:.1f} (target {targets['P_E']})")
    print(f"  CDE={welfare['CDE']:.2f} (target {targets['CDE']})")
    print(f"  ABE={welfare['ABE']:.2f} (target {targets['ABE']})")
    print(f"  EGH={welfare['EGH']:.2f} (target {targets['EGH']})")


def check_assertions(ols_est, gmm_est, welfare, targets, tolerance):
    def check(name, computed, target):
        assert np.isclose(computed, target, rtol=tolerance), (
            f'FAIL {name}: computed {computed:.5g}, target {target:.5g} (rtol={tolerance})'
        )

    for key in PARAMETERS:
        check(f"{key} (OLS)", ols_est[key], targets[key])
        check(f"{key} (GMM)", gmm_est[key], targets[key])

    for key in ('Q_C', 'P_C', 'Q_E', 'P_E', 'CDE', 'ABE', 'EGH'):
        check(key, welfare[key], targets[key])

    assert np.allclose(
        [gmm_est['demand_intercept'], gmm_est['demand_slope']],
        [ols_est['demand_intercept'], ols_est['demand_slope']], rtol=1e-5
    ), 'GMM demand estimates diverge from OLS'
    assert np.allclose(
        [gmm_est['avg_cost_intercept'], gmm_est['avg_cost_slope']],
        [ols_est['avg_cost_intercept'], ols_est['avg_cost_slope']], rtol=1e-5
    ), 'GMM cost estimates diverge from OLS'


if __name__ == '__main__':
    Main()
